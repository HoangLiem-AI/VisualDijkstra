# tao-report.py — Dựng BẢN NHÁP report (Word) từ DEVLOG, số liệu và trace thật.
# Cách dùng (thư mục gốc dự án):
#   node 4_Report/tools/xuat-trace.mjs     # trace thật -> so-lieu/trace-vi-du.json
#   python 4_Report/tools/ve-bieu-do.py    # biểu đồ benchmark -> hinh/bd-*.png
#   python 4_Report/tools/tao-report.py    # -> 4_Report/BaoCao.docx
# Script KHÔNG ghi đè BaoCao.docx đã có (vì SV sửa trực tiếp trên file Word);
# muốn tạo lại từ đầu thì thêm --ghi-de.
# Chỗ tô vàng [[...]] là phần SV phải tự điền / tự viết.

import io
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')
REPORT = Path(__file__).resolve().parents[1]
HINH = REPORT / 'hinh'
OUT = REPORT / 'BaoCao.docx'
if OUT.exists() and '--ghi-de' not in sys.argv:
    sys.exit(f'{OUT} đã tồn tại (có thể SV đã sửa). Thêm --ghi-de nếu muốn tạo lại từ đầu.')

TRACE = json.load(open(REPORT / 'so-lieu' / 'trace-vi-du.json', encoding='utf-8'))
FONT = 'Times New Roman'
MONO = 'Consolas'

doc = Document()

# ---------------- Trang và kiểu chữ ----------------
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
sec.left_margin, sec.right_margin = Cm(3), Cm(2)
TEXT_W = 16  # cm


def set_font(run_or_style, name, size=None, bold=None, color=None):
    f = run_or_style.font
    f.name = name
    rpr = run_or_style.element.get_or_add_rPr() if hasattr(run_or_style.element, 'get_or_add_rPr') else run_or_style.element.rPr
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        fonts.set(qn(k), name)
    # bỏ font theo theme (Calibri Light của Heading) để dùng đúng font chỉ định
    for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if fonts.get(qn(k)) is not None:
            del fonts.attrib[qn(k)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color


normal = doc.styles['Normal']
set_font(normal, FONT, 13)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.3

for name, size in (('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 13)):
    st = doc.styles[name]
    set_font(st, FONT, size, bold=True, color=RGBColor(0, 0, 0))
    st.font.italic = name == 'Heading 3'
    st.paragraph_format.space_before = Pt(12 if name != 'Heading 1' else 0)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
for name in ('List Bullet', 'List Number'):
    set_font(doc.styles[name], FONT, 13)
    doc.styles[name].paragraph_format.space_after = Pt(3)

# ---------------- Hàm tiện ích ----------------
INLINE = re.compile(r'(\*\*.+?\*\*|`.+?`|\[\[.+?\]\]|\*.+?\*)')


def add_runs(p, text, size=None):
    """Hỗ trợ **đậm**, *nghiêng*, `code`, [[chỗ SV điền]] (tô vàng)."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`'):
            r = p.add_run(part[1:-1]); set_font(r, MONO, (size or 13) - 2)
            continue
        elif part.startswith('[['):
            r = p.add_run('[' + part[2:-2] + ']'); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part.startswith('*') and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        if size:
            r.font.size = Pt(size)
    return p


def P(text, align='justify', size=None, indent=True):
    p = doc.add_paragraph()
    p.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'left': WD_ALIGN_PARAGRAPH.LEFT}[align]
    if indent and align == 'justify':
        p.paragraph_format.first_line_indent = Cm(1)
    return add_runs(p, text, size)


def B(text, level=0, numbered=False):
    p = doc.add_paragraph(style='List Number' if numbered else 'List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if level:
        p.paragraph_format.left_indent = Cm(0.63 * (level + 1))
    return add_runs(p, text)


chapter = 0
fig_no = 0
tab_no = 0


def H1(text):
    global chapter, fig_no, tab_no
    chapter += 1
    fig_no = tab_no = 0
    p = doc.add_heading(f'Chương {chapter}. {text}', level=1)
    p.paragraph_format.page_break_before = True
    return p


def H1_plain(text):
    p = doc.add_heading(text, level=1)
    p.paragraph_format.page_break_before = True
    return p


def H2(text):
    return doc.add_heading(text, level=2)


def H3(text):
    return doc.add_heading(text, level=3)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def caption(kind, text):
    global fig_no, tab_no
    if kind == 'Hình':
        fig_no += 1; num = fig_no
    else:
        tab_no += 1; num = tab_no
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    if kind == 'Bảng':
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f'{kind} {chapter}.{num}. '); r.bold = True; r.font.size = Pt(12)
    r = p.add_run(text); r.italic = True; r.font.size = Pt(12)
    return f'{kind} {chapter}.{num}'


def T(headers, rows, widths=None, size=11, title=None, center_cols=()):
    """Bảng có tiêu đề (caption đặt phía trên)."""
    ref = caption('Bảng', title) if title else None
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        add_runs(c.paragraphs[0], f'**{h}**', size)
        shade(c, 'D9E2F3')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            add_runs(p, str(v), size)
            if i in center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.first_line_indent = Cm(0)
        # lặp lại dòng tiêu đề khi bảng sang trang
    trPr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
    if widths:
        t.autofit = False
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
            gc.set(qn('w:w'), str(int(widths[i] * 567)))
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return ref


def CODE(text, size=9.5):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    c = t.rows[0].cells[0]
    shade(c, 'F3F4F6')
    lines = text.strip('\n').split('\n')
    c.text = ''
    for i, line in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(line if line else ' ')
        set_font(r, MONO, size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def FIG(name, text, width=TEXT_W, crop=None):
    img = Image.open(HINH / name)
    if crop:
        img = img.crop(crop)
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(buf, width=Cm(width))
    return caption('Hình', text)


def field(paragraph, code):
    r = paragraph.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
    r = paragraph.add_run()
    i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve'); i.text = code; r._r.append(i)
    r = paragraph.add_run()
    s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate'); r._r.append(s)
    paragraph.add_run('(Nhấn chuột phải → Update Field để cập nhật)')
    r = paragraph.add_run()
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)


# Vùng cắt ảnh chụp (ảnh gốc 3000 × 2640 px, hệ số 2)
CROP_FULL = (0, 0, 3000, 1960)
CROP_MAIN = (645, 120, 2960, 1700)     # đồ thị + điều khiển + pseudo-code + hàng đợi + bảng dist
CROP_LEFT = (0, 120, 640, 1450)

# =====================================================================
# BÌA
# =====================================================================


def center(text, size=13, bold=False, space=6, upper=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space)
    add_runs(p, f'**{text}**' if bold else text, size)
    return p


center('ĐẠI HỌC QUỐC GIA THÀNH PHỐ HỒ CHÍ MINH', 13, True, 0)
center('TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN', 13, True, 0)
center('KHOA [[tên khoa — SV kiểm tra]]', 13, True, 40)
center('BÁO CÁO ĐỒ ÁN CUỐI KỲ', 18, True, 6)
center('MÔN: PHÂN TÍCH VÀ THIẾT KẾ THUẬT TOÁN (CS112)', 14, True, 40)
center('ĐỀ TÀI', 14, True, 4)
center('MÔ PHỎNG TRỰC QUAN TỪNG BƯỚC THUẬT TOÁN DIJKSTRA', 20, True, 4)
center('so sánh cấu trúc dữ liệu (min-heap, mảng) và đối chiếu với Bellman-Ford', 14, False, 60)
for label, val in (('Sinh viên thực hiện', '[[Họ và tên]]'), ('MSSV', '[[MSSV]]'), ('Lớp', '[[CS112.xxxx]]'),
                   ('Giảng viên hướng dẫn', '[[Tên giảng viên]]'),
                   ('Mã nguồn', 'https://github.com/HoangLiem-AI/VisualDijkstra'),
                   ('Video demo', '[[link video]]')):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(3)
    p.paragraph_format.space_after = Pt(2)
    add_runs(p, f'**{label}:** {val}')
p = center('Thành phố Hồ Chí Minh, tháng 10 năm 2026', 13, False, 0)
p.paragraph_format.space_before = Pt(80)

# =====================================================================
# MỤC LỤC
# =====================================================================
p = doc.add_paragraph()
p.paragraph_format.page_break_before = True
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_runs(p, '**MỤC LỤC**', 16)
field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')

# Ghi chú về bản nháp
H1_plain('Ghi chú về bản nháp này')
P('Bản nháp do Claude Code (model Claude Opus 5.5) soạn ngày 18/10/2026 theo yêu cầu của sinh viên, '
  'từ `DEVLOG.md`, mã nguồn, số liệu kiểm thử/benchmark và trace thật do chương trình sinh ra '
  '(không chép tay). Quy trình này được ghi ở mục #14 của DEVLOG.')
P('Các chỗ **tô vàng** dạng [[...]] là phần sinh viên phải tự điền hoặc tự viết, đặc biệt là phần '
  '**kiểm chứng và nhận xét của sinh viên** ở Chương 8. Sau khi sửa xong: xóa trang ghi chú này, '
  'nhấn chuột phải vào mục lục → *Update Field* → *Update entire table*, rồi xuất PDF '
  '(*File → Save As → PDF*).')

# =====================================================================
# CHƯƠNG 1
# =====================================================================
H1('Giới thiệu')
H2('1.1. Bài toán đường đi ngắn nhất một nguồn')
P('Cho đồ thị có hướng có trọng số G = (V, E), hàm trọng số w: E → ℝ và một đỉnh nguồn s. '
  'Độ dài của đường đi p = ⟨v₀, v₁, …, v_k⟩ là tổng trọng số các cạnh trên đường đi. '
  'Khoảng cách ngắn nhất δ(s, v) là giá trị nhỏ nhất của độ dài các đường đi từ s tới v '
  '(bằng ∞ nếu không có đường đi). Bài toán yêu cầu tính δ(s, v) cho mọi đỉnh v, và nếu có đỉnh đích t '
  'thì chỉ ra một đường đi ngắn nhất từ s tới t.')
P('Đây là bài toán nền tảng trong định tuyến mạng, bản đồ, lập lịch vận chuyển. Ví dụ trực tiếp với '
  'sinh viên là tìm hành trình nhanh nhất cho bưu gửi giữa các bưu cục — hướng phát triển được nêu ở Chương 9.')

H2('1.2. Mục tiêu đồ án')
B('Xây dựng chương trình có giao diện cho phép **nhập input → theo dõi từng bước thuật toán → xem output**.')
B('Mô phỏng thuật toán Dijkstra dùng min-heap: tô sáng dòng pseudo-code đang chạy, trạng thái hàng đợi ưu tiên, '
  'bảng dist/prev và đồ thị tại mỗi bước; có thể tiến, lùi, tự chạy, nhảy tới bước bất kỳ.')
B('So sánh cách cài đặt (min-heap và mảng) và chiến lược thiết kế (tham lam – Dijkstra, quy hoạch động – Bellman-Ford), '
  'kể cả trên đồ thị có cạnh âm nơi Dijkstra cho kết quả sai.')
B('Chứng minh tính đúng bằng bộ testcase có đáp án độc lập, và đo thời gian thực tế để đối chiếu với độ phức tạp lý thuyết.')

H2('1.3. Phạm vi')
T(['Mức', 'Nội dung'], [
    ['Bắt buộc', 'Dijkstra dùng min-heap (lazy deletion), mô phỏng từng bước.'],
    ['Mở rộng', 'Dijkstra dùng mảng O(V²); Bellman-Ford (quy hoạch động); chế độ minh họa cạnh âm '
                '(Dijkstra sai, Bellman-Ford đúng); phát hiện và tô chu trình âm.'],
    ['Không làm', 'A*, bản đồ thật, hành trình đường thư theo lịch chuyến (để ở Hướng phát triển).'],
], widths=[3, 13], title='Phạm vi đồ án')

H2('1.4. Sản phẩm nộp')
T(['Thư mục', 'Nội dung'], [
    ['`1_SourceCode/`', 'Mã nguồn: giao diện web (HTML, CSS, JavaScript ES modules, Cytoscape.js), lõi thuật toán, script Node.'],
    ['`2_VideoDemo/`', 'Video demo và kịch bản quay `kich-ban-demo.md`.'],
    ['`3_Testcase/`', '44 test chia 5 nhóm, mỗi test gồm `.in` và `.out`; `README.md` mô tả từng test.'],
    ['`4_Report/`', 'Báo cáo (Word + PDF), hình, số liệu kiểm thử và benchmark.'],
    ['`DEVLOG.md`', 'Nhật ký phát triển: prompt nguyên văn, kết quả AI, lỗi, quyết định thiết kế.'],
], widths=[3.5, 12.5], title='Cấu trúc sản phẩm nộp')
P('Yêu cầu **thống nhất** được đảm bảo bằng một đồ thị duy nhất 6 đỉnh, 11 cạnh: nó là ví dụ ở Chương 5 của report, '
  'là ví dụ được quay trong video, và là testcase `3_Testcase/nho/01-vi-du-6-dinh`. Đồ thị này được khai báo một lần '
  'trong `src/core/samples.js` và mọi nơi khác đều lấy từ đó.')

H2('1.5. Bố cục báo cáo')
P('Chương 2 trình bày cơ sở lý thuyết. Chương 3 và 4 trình bày thiết kế và cài đặt. Chương 5 chạy từng bước trên '
  'ví dụ thống nhất. Chương 6 và 7 là kiểm thử và đánh giá hiệu năng. Chương 8 tường thuật quá trình xây dựng, '
  'trong đó có cách sử dụng AI. Chương 9 là kết luận và hướng phát triển.')

# =====================================================================
# CHƯƠNG 2
# =====================================================================
H1('Cơ sở lý thuyết')
H2('2.1. Ba chiến lược giải cùng một bài toán')
T(['Chiến lược', 'Thuật toán', 'Ý tưởng', 'Độ phức tạp'], [
    ['Vét cạn', 'Duyệt mọi đường đi đơn (DFS)', 'Thử hết, lấy đường ngắn nhất', 'O(n!) — chỉ dùng làm đáp án chuẩn với n ≤ 8'],
    ['Tham lam', 'Dijkstra', 'Mỗi lần chốt đỉnh có dist tạm thời nhỏ nhất', 'O((V + E) log V) với heap; O(V² + E) với mảng'],
    ['Quy hoạch động', 'Bellman-Ford', 'Tính dần đường đi ngắn nhất dùng ≤ k cạnh', 'O(V · E)'],
], widths=[2.6, 3.4, 5.4, 4.6], title='Các chiến lược thiết kế được dùng trong đồ án')
P('Cả ba đều dựa trên **cấu trúc con tối ưu**: một đoạn con của đường đi ngắn nhất cũng là đường đi ngắn nhất. '
  'Phép toán cơ bản chung là **relax** cạnh (u, v): nếu dist[u] + w(u, v) < dist[v] thì cập nhật dist[v] và prev[v] = u. '
  'Các thuật toán chỉ khác nhau ở *thứ tự* relax các cạnh.')

H2('2.2. Thuật toán Dijkstra')
H3('Ý tưởng')
P('Dijkstra duy trì tập S các đỉnh đã **chốt** (dist đã là khoảng cách ngắn nhất). Mỗi bước chọn đỉnh u ∉ S có dist[u] '
  'nhỏ nhất, đưa u vào S, rồi relax mọi cạnh ra khỏi u. Đây là lựa chọn tham lam: đỉnh gần nhất trong số chưa chốt '
  'được coi là đã biết khoảng cách chính xác.')
H3('Pseudo-code (bản min-heap, đúng như hiển thị trên giao diện)')
CODE('\n'.join(f'{i + 1:>2}  {line}' for i, line in enumerate([
    'for mỗi đỉnh v: dist[v] ← ∞; prev[v] ← null',
    'dist[s] ← 0; PQ.push(s, 0)',
    'while PQ khác rỗng:',
    '    (u, d) ← PQ.popMin()',
    '    if u đã chốt: continue          ▷ phần tử cũ (lazy deletion)',
    '    chốt u                          ▷ dist[u] là ngắn nhất',
    '    for mỗi cạnh (u, v, w) với v chưa chốt:',
    '        nd ← relax((u, v, w), dist[u])   ▷ mặc định dist[u] + w',
    '        if nd < dist[v]:',
    '            dist[v] ← nd; prev[v] ← u; PQ.push(v, nd)',
    'return dist, prev'])))
H3('Tính đúng')
P('**Mệnh đề.** Nếu mọi trọng số không âm thì khi một đỉnh u được chốt, dist[u] = δ(s, u).')
P('**Chứng minh** (phản chứng). Giả sử u là đỉnh *đầu tiên* được chốt mà dist[u] > δ(s, u). Khi đó u ≠ s và tồn tại '
  'đường đi ngắn nhất P từ s tới u. Gọi y là đỉnh đầu tiên trên P chưa được chốt tại thời điểm chốt u, và x là đỉnh '
  'đứng ngay trước y trên P (x đã chốt). Vì x được chốt trước u nên dist[x] = δ(s, x), và khi chốt x cạnh (x, y) đã được '
  'relax, do đó dist[y] ≤ δ(s, x) + w(x, y) = δ(s, y). Vì y nằm trên P trước u và **các trọng số không âm**, '
  'δ(s, y) ≤ δ(s, u). Suy ra dist[y] ≤ δ(s, u) < dist[u], trái với việc u được chọn vì có dist nhỏ nhất. ∎')
P('Chỗ dùng giả thiết "trọng số không âm" chính là chỗ Dijkstra hỏng khi có cạnh âm (mục 2.5).')
H3('Độ phức tạp')
B('**Bản mảng:** V lần tìm đỉnh nhỏ nhất, mỗi lần duyệt mảng O(V); mỗi cạnh relax một lần. Tổng **O(V² + E)**.')
B('**Bản min-heap với lazy deletion:** mỗi lần relax thành công đẩy một phần tử mới vào heap (không có decrease-key), '
  'nên heap có tối đa E + 1 phần tử. Mỗi push/pop tốn O(log E) = O(log V) (vì E ≤ V²). Tổng **O((V + E) log V)**, '
  'bộ nhớ O(V + E).')
B('Đồ thị thưa (E ≈ V): bản heap O(V log V) tốt hơn hẳn O(V²). Đồ thị dày (E ≈ V²): bản heap O(V² log V), '
  'bản mảng O(V²) lại tốt hơn. Chương 7 kiểm chứng điều này bằng số đo.')

H2('2.3. Min-heap nhị phân và lazy deletion')
P('Min-heap nhị phân là cây nhị phân đầy đủ lưu trong mảng: nút i có con 2i + 1, 2i + 2 và cha ⌊(i − 1)/2⌋; '
  'khóa của cha luôn ≤ khóa của con nên phần tử nhỏ nhất ở gốc. `push` đặt phần tử vào cuối rồi đẩy lên (sift-up), '
  '`pop` đưa phần tử cuối lên gốc rồi đẩy xuống (sift-down); cả hai O(log n).')
P('Dijkstra cần giảm khóa của một đỉnh khi relax thành công (decrease-key). Thay vì cài decrease-key (phải lưu vị trí '
  'của mỗi đỉnh trong heap), đồ án dùng **lazy deletion**: đẩy luôn cặp (v, dist mới) vào heap; cặp cũ vẫn nằm lại và '
  'sẽ bị **bỏ qua khi lấy ra** nếu đỉnh đã chốt. Trên giao diện, bước bỏ qua này có loại `skip-stale`, và phần tử cũ '
  'trong hàng đợi được gạch ngang để người xem thấy rõ.')
P('Khi hai khóa bằng nhau, heap ưu tiên đỉnh có số nhỏ hơn. Nhờ vậy trace luôn tái tạo được và bản heap, bản mảng '
  'chốt các đỉnh theo cùng thứ tự.')

H2('2.4. Bellman-Ford (quy hoạch động)')
P('Gọi d_k(v) là độ dài đường đi ngắn nhất từ s tới v dùng **tối đa k cạnh**. Công thức truy hồi:')
P('d_k(v) = min( d_{k−1}(v), min_{(u,v)∈E} [ d_{k−1}(u) + w(u, v) ] ),   d_0(s) = 0, d_0(v) = ∞ với v ≠ s.', align='center')
P('Nếu không có chu trình âm, đường đi ngắn nhất là đường đi đơn, dùng tối đa V − 1 cạnh, nên sau V − 1 vòng '
  '(mỗi vòng relax mọi cạnh) ta có kết quả đúng. Thêm một vòng nữa: nếu vẫn còn cạnh relax được thì tồn tại chu trình '
  'âm tới được từ s và bài toán không có lời giải; chương trình lần theo prev để chỉ ra chu trình đó. '
  'Độ phức tạp O(V · E). Cài đặt có **dừng sớm**: nếu một vòng không cập nhật gì thì các vòng sau cũng không, nên dừng.')
P('Có thể xem Dijkstra cũng là quy hoạch động trên cùng cấu trúc con tối ưu, nhưng thứ tự giải các bài toán con '
  'được chọn bằng tham lam (đỉnh gần nhất trước), nhờ đó mỗi cạnh chỉ cần relax một lần.')

H2('2.5. Vì sao Dijkstra sai khi có cạnh âm')
P('Xét đồ thị 4 đỉnh: 1→2 (2), 1→3 (5), 3→2 (−4), 2→4 (3), nguồn s = 1. Dijkstra chốt đỉnh 2 với dist = 2 ngay sau '
  'đỉnh 1, vì 2 là nhỏ nhất. Nhưng đường 1→3→2 dài 5 + (−4) = 1 < 2. Khi xét đỉnh 3 thì đỉnh 2 đã chốt, cạnh 3→2 bị bỏ '
  'qua, và đỉnh 4 nhận dist = 5 thay vì 4. Kết quả Dijkstra: `0 2 5 5`; đúng (Bellman-Ford): `0 1 5 4`. '
  'Ví dụ này có sẵn trên giao diện (mục 5.5) và là testcase `dac-biet/01-canh-am-dijkstra-sai`.')
P('Để minh họa được điều này, cài đặt giữ đúng Dijkstra kinh điển: kiểm tra phần tử cũ bằng "u đã chốt" và '
  '**không relax tới đỉnh đã chốt**. Nếu cho phép xử lý lại đỉnh, thuật toán trở thành dạng label-correcting và có thể '
  'tình cờ cho kết quả đúng, làm mất ý nghĩa minh họa. Ở chế độ bình thường, Dijkstra từ chối chạy và báo lỗi rõ ràng '
  'khi đồ thị có cạnh âm.')

H2('2.6. So sánh tổng hợp')
T(['Tiêu chí', 'Dijkstra (min-heap)', 'Dijkstra (mảng)', 'Bellman-Ford'], [
    ['Chiến lược', 'Tham lam', 'Tham lam', 'Quy hoạch động'],
    ['Thời gian', 'O((V + E) log V)', 'O(V² + E)', 'O(V · E)'],
    ['Bộ nhớ', 'O(V + E)', 'O(V)', 'O(V)'],
    ['Cạnh âm', 'Không (báo lỗi)', 'Không (báo lỗi)', 'Có'],
    ['Chu trình âm', '—', '—', 'Phát hiện và chỉ ra'],
    ['Phù hợp', 'Đồ thị thưa', 'Đồ thị dày', 'Có cạnh âm'],
], widths=[3, 4.4, 4.2, 4.4], title='So sánh ba thuật toán')

# =====================================================================
# CHƯƠNG 3
# =====================================================================
H1('Phân tích và thiết kế')
H2('3.1. Yêu cầu chức năng')
B('Nhập đồ thị dạng văn bản (kiểu Wecode), chọn có hướng / vô hướng, chọn thuật toán, chọn ví dụ có sẵn.')
B('Phát lại từng bước: Tiến, Lùi, Về đầu, Tới cuối, Tự chạy với tốc độ điều chỉnh được, thanh trượt nhảy tới bước bất kỳ, phím tắt.')
B('Mỗi bước hiển thị: đồ thị tô màu theo trạng thái, dòng pseudo-code đang chạy, hàng đợi ưu tiên, bảng dist/prev, câu giải thích tiếng Việt, nhật ký.')
B('Output đúng định dạng đề; input sai, cạnh âm (với Dijkstra), chu trình âm được báo lỗi rõ ràng, không crash.')
B('Đồ thị lớn (tới 10⁵ đỉnh, 2·10⁵ cạnh) chạy bằng script, không vẽ.')

H2('3.2. Kiến trúc: lõi sinh trace, giao diện chỉ phát lại')
P('Quyết định quan trọng nhất là **tách lõi thuật toán khỏi giao diện**. Thuật toán chạy một lần, trả về '
  '`{ result, trace }`, trong đó `trace` là mảng các bước. Giao diện chỉ phát lại trace, không tự tính lại. Nhờ vậy:')
B('Cùng một mã thuật toán được giao diện web, script kiểm thử và script benchmark dùng chung; cái được kiểm thử chính là cái được mô phỏng.')
B('Lùi bước trở nên đơn giản (chỉ đổi chỉ số trong mảng trace), không phải "chạy ngược" thuật toán.')
B('Lõi không đụng tới DOM nên chạy được bằng Node.js.')
T(['Tầng', 'File', 'Vai trò'], [
    ['Lõi (`src/core/`)', '`parser.js`', 'Đọc và kiểm tra input, báo lỗi kèm số dòng; định dạng output.'],
    ['', '`graph.js`', 'Danh sách kề, đường đi từ prev, hàm `relax` mặc định, kiểm tra cạnh âm.'],
    ['', '`min-heap.js`', 'Heap nhị phân, khóa bằng nhau thì đỉnh nhỏ trước.'],
    ['', '`dijkstra-heap.js`, `dijkstra-array.js`, `bellman-ford.js`', 'Ba thuật toán, cùng trả về `{ result, trace }`.'],
    ['', '`algorithms.js`, `samples.js`, `solve.js`', 'Danh mục thuật toán; đồ thị ví dụ dùng chung; chuỗi input → chuỗi output.'],
    ['Giao diện (`src/ui/`)', '`renderer.js`', 'Vẽ đồ thị bằng Cytoscape.js, tô màu theo bước.'],
    ['', '`player.js`', 'Tiến / Lùi / Tự chạy (không đụng DOM).'],
    ['', '`panels.js`, `app.js`', 'Pseudo-code, hàng đợi, bảng dist, nhật ký; nối các phần.'],
    ['Script (`scripts/`)', '`gen-tests.js`, `run-tests.js`, `brute-force.js`, `benchmark.js`', 'Sinh testcase có seed, chấm, đáp án vét cạn/Floyd, đo thời gian.'],
], widths=[3.4, 5.6, 7], size=10.5, title='Các module của chương trình')

H2('3.3. Cấu trúc một bước trace')
T(['Trường', 'Ý nghĩa'], [
    ['`type`', '`init`, `extract`, `skip-stale`, `relax-ok`, `relax-fail`, `done`; Bellman-Ford thêm `round`, `negative-cycle`.'],
    ['`line`', 'Số dòng pseudo-code đang chạy, để tô sáng.'],
    ['`current`, `edge`', 'Đỉnh đang xét và cạnh đang relax.'],
    ['`dist`, `prev`, `finalized`, `pq`', '**Bản sao** tại thời điểm đó (không phải tham chiếu), nên lùi bước luôn đúng.'],
    ['`message`', 'Câu giải thích tiếng Việt, ví dụ "Relax 3→2: 4 → 3".'],
    ['`round`, `maxRound`, `cycle`', 'Chỉ Bellman-Ford: vòng hiện tại và chu trình âm.'],
], widths=[4.6, 11.4], title='Các trường của một bước trace')
P('Vì mỗi bước sao chép mảng O(V), trace tốn O(số bước × V) bộ nhớ. Điều này chấp nhận được với đồ thị ≤ 30 đỉnh '
  'trên giao diện; khi chạy đồ thị lớn, script truyền `recordTrace: false` để tắt ghi trace.')

H2('3.4. Hàm relax(edge, t)')
P('Mọi thuật toán tính giá trị mới qua `relax(edge, t)` (mặc định trả về `t + edge.w`) thay vì cộng trực tiếp trong vòng '
  'lặp. Lý do: hướng phát triển là bài toán hành trình đường thư, trong đó cạnh là các chuyến có giờ đi/giờ đến; '
  'khi đó chỉ cần thay hàm relax bằng "giờ đến sớm nhất nếu có mặt ở u lúc t" mà không sửa thuật toán.')

H2('3.5. Định dạng input/output và xử lý lỗi')
CODE('n m s [t]          ▷ số đỉnh, số cạnh, đỉnh nguồn, đỉnh đích (không bắt buộc)\n'
     'u1 v1 w1          ▷ m dòng cạnh, đỉnh đánh số 1..n\n...\num vm wm')
B('Ràng buộc: 1 ≤ n ≤ 10⁵; 0 ≤ m ≤ 2·10⁵; |w| ≤ 10⁹ (Dijkstra yêu cầu w ≥ 0). Tổng lớn nhất khoảng 10¹⁴ < 2⁵³ nên kiểu `Number` của JavaScript vẫn chính xác.')
B('Output: n dòng dist (−1 nếu không tới được); nếu có t thì thêm dòng đường đi `s -> ... -> t` (−1 nếu không có đường).')
B('Chu trình âm (Bellman-Ford): một dòng `Phát hiện chu trình âm`. Cạnh âm với Dijkstra: thông báo lỗi kèm gợi ý dùng Bellman-Ford.')
B('Input sai: thông báo kèm số dòng, ví dụ `Dòng 4: "x" không phải số nguyên (w)` (Hình 4.2).')

H2('3.6. Các quyết định thiết kế')
P('Bảng dưới lấy từ mục "Các quyết định thiết kế" trong DEVLOG. Các quyết định do AI đề xuất được đánh dấu để sinh viên xác nhận.')
T(['Quyết định', 'Phương án đã cân nhắc', 'Lý do'], [
    ['Phạm vi: Dijkstra heap + mảng + Bellman-Ford', 'Chỉ Dijkstra; thêm A* và bản đồ thật', 'Cân bằng sáng tạo và thời gian (hạn 18/10)'],
    ['Web tĩnh + Cytoscape.js; script bằng Node', 'Python (Streamlit/PyQt); .NET', 'Vẽ đồ thị tốt, chạy trên trình duyệt, dùng chung lõi'],
    ['Lazy deletion, kiểm tra "u đã chốt"; không relax tới đỉnh đã chốt', 'Kiểm tra d > dist[u] và cho xử lý lại đỉnh', 'Giữ đúng Dijkstra kinh điển để minh họa sai với cạnh âm'],
    ['Parser cho phép w âm; Dijkstra tự báo lỗi', 'Parser cấm cạnh âm', 'Dùng chung parser cho Bellman-Ford'],
    ['Đồ thị ví dụ để trong `samples.js`', 'Mỗi nơi tự chép đồ thị', 'Report, video, testcase dùng đúng một đồ thị'],
    ['Bellman-Ford duyệt cạnh theo thứ tự input, dừng sớm', 'Luôn chạy đủ n − 1 vòng', 'Trace khớp file input, ngắn hơn khi demo'],
    ['Chấm đường đi bằng kiểm tra hợp lệ + độ dài', 'So khớp chuỗi', 'Có thể có nhiều đường ngắn nhất'],
    ['Đáp án test lớn bằng Bellman-Ford', 'Dùng Dijkstra làm đáp án', 'Nguồn đáp án độc lập với thuật toán được kiểm tra'],
], widths=[5.6, 5, 5.4], size=10.5, title='Các quyết định thiết kế chính')

# =====================================================================
# CHƯƠNG 4
# =====================================================================
H1('Cài đặt và giao diện')
H2('4.1. Vòng lặp chính của Dijkstra (min-heap)')
P('Đoạn mã dưới trích từ `src/core/dijkstra-heap.js` (đã lược phần đếm thống kê). Mỗi nhánh gọi `record(...)` để ghi một bước trace với số dòng pseudo-code tương ứng.')
CODE('''while (!pq.isEmpty()) {
  // Dòng 4: lấy phần tử có khóa nhỏ nhất
  const { vertex: u, key: d } = pq.pop();
  // Dòng 5: lazy deletion — u đã chốt thì (u, d) là bản cũ, bỏ qua
  if (finalized[u]) {
    record('skip-stale', 5, u, null, `Lấy (${u}, ${d}) nhưng đỉnh ${u} đã chốt ...`);
    continue;
  }
  // Dòng 6: chốt u, dist[u] không thay đổi nữa
  finalized[u] = true;
  record('extract', 6, u, null, `Lấy (${u}, ${d}) ra khỏi hàng đợi, chốt dist[${u}] = ${d}`);
  // Dòng 7: duyệt các cạnh ra khỏi u
  for (const edge of graph.adj[u]) {
    const v = edge.v;
    if (finalized[v]) { record('relax-fail', 7, u, edge, `Bỏ qua ${u}→${v}: đỉnh ${v} đã chốt`); continue; }
    // Dòng 8: tính giá trị mới qua hàm relax (không cộng trực tiếp)
    const nd = relax(edge, dist[u]);
    // Dòng 9–10: cập nhật nếu tốt hơn, đẩy bản mới vào heap
    if (nd < dist[v]) {
      dist[v] = nd; prev[v] = u; pq.push(v, nd);
      record('relax-ok', 10, u, edge, `Relax ${u}→${v}: ${old} → ${nd}`);
    } else {
      record('relax-fail', 9, u, edge, `Relax ${u}→${v}: ... ≥ ..., giữ nguyên`);
    }
  }
}''')
H2('4.2. Min-heap: thao tác sift-down')
CODE('''#siftDown(i) {
  const n = this.items.length;
  for (;;) {
    const l = 2 * i + 1;
    const r = l + 1;
    let smallest = i;
    if (l < n && this.#less(l, smallest)) smallest = l;
    if (r < n && this.#less(r, smallest)) smallest = r;
    if (smallest === i) break;
    this.#swap(i, smallest);
    i = smallest;
  }
}''')
P('Dòng so sánh với con phải ở trên được dùng làm "lỗi cài cắm" khi kiểm tra chất lượng bộ test (mục 6.5).')

H2('4.3. Giao diện')
P('Giao diện gồm ba cột (Hình 4.1). Cột trái: chọn thuật toán, ví dụ có sẵn, ô nhập input, tùy chọn vô hướng và minh họa '
  'cạnh âm, output và nhật ký các bước (bấm vào một dòng để nhảy tới bước đó). Cột giữa: đồ thị, câu giải thích của bước '
  'hiện tại, các nút điều khiển và chú thích màu. Cột phải: pseudo-code với dòng đang chạy được tô sáng, hàng đợi ưu tiên '
  '(Dijkstra mảng: danh sách đỉnh ứng viên; Bellman-Ford: thanh tiến độ vòng lặp) và bảng dist/prev.')
FIG('h01-khoi-tao.png', 'Giao diện tại bước 1 (khởi tạo) của ví dụ thống nhất', crop=CROP_FULL)
T(['Màu / ký hiệu', 'Ý nghĩa'], [
    ['Đỉnh xám', 'Chưa thăm (dist = ∞)'],
    ['Viền xanh dương', 'Tạm thời: dist < ∞, chưa chốt'],
    ['Viền cam', 'Đỉnh đang xét'],
    ['Nền xanh lá', 'Đã chốt'],
    ['Cạnh xanh lá / đỏ nét đứt', 'Relax thành công / không cải thiện'],
    ['Cạnh xanh dương', 'Cây prev (cạnh dẫn tới dist hiện tại)'],
    ['Cạnh tím (bước cuối)', 'Đường đi ngắn nhất s → t'],
    ['Cạnh và đỉnh đỏ', 'Chu trình âm'],
    ['Phần tử gạch ngang trong hàng đợi', 'Phần tử cũ, sẽ bị bỏ qua (lazy deletion)'],
], widths=[6, 10], title='Ý nghĩa màu trên đồ thị')
P('Phím tắt: ← → lùi/tiến, Space tự chạy, Home/End về đầu/cuối, Ctrl + Enter chạy thuật toán. Giao diện giới hạn '
  '30 đỉnh, 150 cạnh để hình vẽ còn đọc được.')
FIG('h10-input-sai.png', 'Input sai được báo lỗi kèm số dòng, không crash', width=6.5, crop=CROP_LEFT)

H2('4.4. Cách chạy')
CODE('cd 1_SourceCode\n'
     'node scripts/serve.js            ▷ mở http://localhost:8080\n'
     'node scripts/demo-trace.js       ▷ in trace ra console\n'
     'node scripts/run-tests.js        ▷ chấm toàn bộ 3_Testcase\n'
     'node scripts/benchmark.js        ▷ đo thời gian')
P('Không mở trực tiếp `index.html` bằng double-click vì trình duyệt chặn ES modules qua `file://`. Chương trình không '
  'có thư viện cần cài (Cytoscape.js tải qua CDN), chỉ cần Node.js. Chi tiết trong `README.md`.')

# =====================================================================
# CHƯƠNG 5
# =====================================================================
H1('Ví dụ minh họa thống nhất')
H2('5.1. Đồ thị ví dụ')
P('Đồ thị có hướng 6 đỉnh, 11 cạnh, nguồn s = 1, đích t = 6. Đồ thị được chọn để trace có **đủ mọi loại bước**: relax '
  'thành công, relax không cải thiện, bỏ qua vì đỉnh đã chốt, và bỏ qua phần tử cũ trong heap. Cạnh `6 4 1` được thêm '
  'vào sau lần chạy đầu chính vì lý do này (lỗi #1 trong DEVLOG).')
CODE(TRACE['heap']['input'])
P('Output (giống nhau ở cả 3 thuật toán; trùng với `3_Testcase/nho/01-vi-du-6-dinh.out`):')
CODE(TRACE['heap']['output'])

H2('5.2. Dijkstra (min-heap): toàn bộ 23 bước')
P('Bảng dưới được sinh tự động từ trace thật (`4_Report/tools/xuat-trace.mjs`), trùng khớp với bộ đếm "Bước k / 23" trên giao diện và trong video.')
TYPE_VN = {'init': 'Khởi tạo', 'extract': 'Lấy ra, chốt', 'skip-stale': 'Bỏ qua bản cũ', 'relax-ok': 'Relax: cập nhật',
           'relax-fail': 'Relax: không đổi', 'done': 'Kết thúc', 'round': 'Vòng', 'negative-cycle': 'Chu trình âm'}
rows = [[s['k'], TYPE_VN[s['type']], s['line'], s['message'], s['pq'] or '∅'] for s in TRACE['heap']['steps']]
T(['#', 'Loại', 'Dòng', 'Diễn giải', 'Hàng đợi sau bước (tăng dần)'], rows,
  widths=[0.8, 2.5, 1.4, 6.6, 4.7], size=9.5, center_cols=(0, 2), title='Trace Dijkstra (min-heap) trên ví dụ thống nhất')
st = TRACE['heap']['stats']
P(f'Thống kê: {st["pushes"]} lần push, {st["pops"]} lần pop, trong đó {st["staleSkips"]} phần tử cũ bị bỏ qua; '
  f'{st["relaxations"]} lần tính relax, {st["improvements"]} lần cập nhật.')

H2('5.3. Các bước đáng chú ý')
H3('Bước 6 — relax làm giảm khoảng cách')
P('Đỉnh 3 vừa được chốt với dist = 1. Relax cạnh 3→2 cho 1 + 2 = 3 < 4, nên dist[2] giảm từ 4 xuống 3, prev[2] = 3, '
  'và cặp (2, 3) được đẩy vào heap. Cặp cũ (2, 4) **không bị xóa**: nó vẫn nằm trong heap và được gạch ngang.')
FIG('h02-heap-buoc-06.png', 'Bước 6: relax 3→2, dist[2] giảm 4 → 3; (2, 4) trở thành phần tử cũ', crop=CROP_MAIN)
H3('Bước 12 — lazy deletion')
P('Heap trả về (2, 4), nhưng đỉnh 2 đã được chốt ở bước 9 với dist = 3, nên phần tử này bị bỏ qua (dòng 5 của pseudo-code). '
  'Đây là cái giá của lazy deletion: heap lớn hơn một chút (10 lần push cho 6 đỉnh), đổi lại không cần decrease-key.')
FIG('h03-heap-buoc-12-skip-stale.png', 'Bước 12: lấy ra phần tử cũ (2, 4) và bỏ qua', crop=CROP_MAIN)
H3('Bước 21 — không relax tới đỉnh đã chốt')
P('Đỉnh 6 được chốt với dist = 13; cạnh 6→4 dẫn tới đỉnh 4 đã chốt nên bị bỏ qua (dòng 7). Với trọng số không âm, '
  'việc này không bao giờ làm mất kết quả đúng (mục 2.2).')
FIG('h04-heap-buoc-21-da-chot.png', 'Bước 21: bỏ qua cạnh 6→4 vì đỉnh 4 đã chốt', crop=CROP_MAIN)
H3('Bước 23 — kết quả')
P('Hàng đợi rỗng, thuật toán kết thúc. Đường đi ngắn nhất 1 → 3 → 2 → 4 → 5 → 6 có độ dài 1 + 2 + 5 + 2 + 3 = 13, được tô tím.')
FIG('h05-heap-buoc-23-ket-qua.png', 'Bước 23: đường đi ngắn nhất 1 → 3 → 2 → 4 → 5 → 6, độ dài 13', crop=CROP_FULL)

H2('5.4. Cùng đồ thị với Dijkstra mảng và Bellman-Ford')
P(f'**Dijkstra (mảng)** cho {len(TRACE["array"]["steps"])} bước: các đỉnh được chốt theo đúng thứ tự như bản heap '
  '(1, 3, 2, 4, 5, 6) nhưng không có bước "bỏ qua phần tử cũ", vì mỗi lần chọn đỉnh là duyệt lại cả mảng dist. '
  'Khung bên phải hiện danh sách đỉnh ứng viên thay cho heap.')
FIG('h06-mang-buoc-06.png', 'Dijkstra (mảng), bước 6: danh sách ứng viên thay cho hàng đợi ưu tiên', crop=CROP_MAIN)
P(f'**Bellman-Ford** cho {len(TRACE["bf"]["steps"])} bước. Vì các cạnh trong input gần như đã theo thứ tự topo, '
  'vòng 1 đã tìm ra mọi khoảng cách đúng; vòng 2 duyệt lại 11 cạnh mà không cập nhật gì nên thuật toán **dừng sớm**, '
  'không cần chạy đủ n − 1 = 5 vòng. Kết quả trùng với Dijkstra.')
FIG('h07-bf-vong-2.png', 'Bellman-Ford, bước 14: bắt đầu vòng 2/5', crop=CROP_MAIN)
T(['Thuật toán', 'Số bước trace', 'Ghi chú'], [
    ['Dijkstra (min-heap)', len(TRACE['heap']['steps']), f'{st["staleSkips"]} bước bỏ qua phần tử cũ'],
    ['Dijkstra (mảng)', len(TRACE['array']['steps']), '6 lần duyệt mảng tìm đỉnh nhỏ nhất'],
    ['Bellman-Ford', len(TRACE['bf']['steps']), f'{TRACE["bf"]["stats"]["rounds"]} vòng, dừng sớm; {TRACE["bf"]["stats"]["relaxations"]} lần relax'],
], widths=[4.5, 3, 8.5], center_cols=(1,), title='So sánh ba thuật toán trên ví dụ thống nhất')

H2('5.5. Cạnh âm và chu trình âm')
P('Với ví dụ ở mục 2.5, khi bật **Minh họa cạnh âm**, Dijkstra vẫn chạy và tới bước cuối hiện bảng so sánh với '
  'Bellman-Ford: sai ở đỉnh 2 (2 thay vì 1) và đỉnh 4 (5 thay vì 4). Bước quyết định là bước 8: "Bỏ qua 3→2: đỉnh 2 đã chốt".')
FIG('h08-canh-am-so-sanh.png', 'Dijkstra sai trên đồ thị có cạnh âm; bảng so sánh với Bellman-Ford', crop=(0, 0, 3000, 2300))
P('Với đồ thị có chu trình 2 → 3 → 4 → 2 tổng 1 + (−3) + 1 = −1, sau n − 1 = 4 vòng cạnh 2→3 vẫn relax được. '
  'Bellman-Ford báo chu trình âm, tô đỏ chu trình và output `Phát hiện chu trình âm`.')
FIG('h09-chu-trinh-am.png', 'Bellman-Ford phát hiện và tô đỏ chu trình âm 2 → 3 → 4 → 2', crop=(0, 0, 3000, 2280))

# =====================================================================
# CHƯƠNG 6
# =====================================================================
H1('Kiểm thử')
H2('6.1. Chiến lược')
B('**Đáp án độc lập với thuật toán được kiểm tra:** vét cạn mọi đường đi đơn (n ≤ 8), Floyd-Warshall (n ≤ 300), Bellman-Ford có dừng sớm (đồ thị lớn). Khi sinh test, các nguồn đáp án được đối chiếu chéo với nhau; lệch là dừng.')
B('**Tái tạo được:** test sinh bằng `scripts/gen-tests.js` với bộ sinh số ngẫu nhiên mulberry32 có seed (seed = 20261018 + số thứ tự test).')
B('**Cả 3 cài đặt** chạy trên mọi test qua cùng hàm `solve` (chuỗi input → chuỗi output) như một chương trình nộp chấm.')
H2('6.2. Bộ testcase')
T(['Nhóm', 'Số test', 'Nội dung'], [
    ['`nho`', 11, 'Ví dụ thống nhất + 10 đồ thị ngẫu nhiên n ≤ 8; đáp án vét cạn.'],
    ['`bien`', 9, 'n = 1; m = 0; s = t; trọng số 0; trọng số 10⁹ (dist tới ~10¹²); khuyên và cạnh song song; không có t; t không tới được; n = 10⁵ không cạnh.'],
    ['`dac-biet`', 9, 'Cạnh âm (Dijkstra phải báo lỗi); chu trình âm tới được / không tới được từ s; chu trình tổng 0; đồ thị đầy đủ; hình sao; lưới nhiều đường bằng nhau; nhiều phần tử cũ trong heap.'],
    ['`lon`', 4, 'n = 10⁵, m = 2·10⁵ ngẫu nhiên; chuỗi 10⁵ đỉnh w = 10⁹ (dist ≈ 10¹⁴); đồ thị dày n = 450; lưới 220 × 220.'],
    ['`khong-hop-le`', 11, 'Input rỗng, thiếu số, chữ, số thực, n = 0, đỉnh ngoài phạm vi, w quá lớn, thiếu/thừa dòng cạnh, m quá lớn.'],
], widths=[2.8, 1.6, 11.6], center_cols=(1,), title='Năm nhóm testcase (44 test)')
H2('6.3. Cách chấm')
B('Các dòng khoảng cách phải khớp chính xác.')
B('Dòng đường đi: chấp nhận mọi đường hợp lệ từ s tới t có độ dài bằng dist[t], vì có thể có nhiều đường ngắn nhất (ví dụ test lưới mọi trọng số 1). So khớp chuỗi sẽ đánh trượt oan.')
B('Test có cạnh âm: Dijkstra (heap, mảng) đạt khi báo đúng lỗi cạnh âm; Bellman-Ford phải ra đáp án.')
B('Nhóm `khong-hop-le`: output phải đúng thông báo lỗi, ví dụ `Lỗi: Dòng 2: "a" không phải số nguyên (w)`.')
H2('6.4. Kết quả')
P('`node scripts/run-tests.js` chạy 44 test × 3 thuật toán: **128 đạt, 0 không đạt, 4 bỏ qua**. Bốn lượt bỏ qua là '
  'Dijkstra mảng trên các test n ≥ 48 400 (O(V²) quá chậm). Bảng đầy đủ ở `4_Report/so-lieu/ket-qua-kiem-thu.md`; '
  'bảng dưới trích nhóm đồ thị lớn.')
T(['Test', 'n / m', 'Heap', 'Mảng', 'Bellman-Ford'], [
    ['lon/01-thua-ngau-nhien', '100 000 / 200 000', '✓ 472.7 ms', '— (n lớn)', '✓ 1578.7 ms'],
    ['lon/02-chuoi-trong-so-max', '100 000 / 199 998', '✓ 308.4 ms', '— (n lớn)', '✓ 361.4 ms'],
    ['lon/03-day', '450 / 200 000', '✓ 231.3 ms', '✓ 214.8 ms', '✓ 332.9 ms'],
    ['lon/04-luoi', '48 400 / 192 720', '✓ 242.4 ms', '— (n lớn)', '✓ 5075.1 ms'],
], widths=[4.6, 3.6, 2.6, 2.4, 2.8], size=10.5, title='Kết quả nhóm đồ thị lớn (thời gian gồm cả đọc input, ghi output)')
P('Ngoài bộ test chính thức, trong lúc cài đặt đã kiểm tra chéo 3000 đồ thị ngẫu nhiên (n ≤ 9) với Floyd-Warshall: 3 thuật '
  'toán cho cùng dist; Bellman-Ford đúng với 387 ca có cạnh âm không chu trình âm và phát hiện đúng 566 ca chu trình âm. 0 ca sai. '
  'Test `lon/02` xác nhận độ chính xác số lớn: dist cuối 99 999 000 000 000 đúng (< 2⁵³).')
H2('6.5. Bộ test có bắt được lỗi không? (mutation test)')
P('Để biết bộ test có thực sự phát hiện lỗi, đã **cố ý cài lỗi**: xóa dòng so sánh với con phải trong `siftDown` của min-heap '
  '(mục 4.2), rồi chạy lại `run-tests.js`. Kết quả: 5 test không đạt, đều ở bản heap — `dac-biet/06-day-du`, '
  '`dac-biet/09-nhieu-phan-tu-cu` và 3 test nhóm `lon`.')
P('Điều đáng chú ý là **cả 11 test nhỏ ngẫu nhiên đều không phát hiện được lỗi này**, vì heap nhỏ ít khi rơi vào trường hợp con phải '
  'nhỏ hơn con trái. Bài học: cần nhóm test đặc biệt được thiết kế có chủ đích (nhiều phần tử cũ trong heap) và test lớn. '
  'Sau đó đã khôi phục mã và chạy lại 128/128 đạt.')

# =====================================================================
# CHƯƠNG 7
# =====================================================================
H1('Đánh giá hiệu năng')
H2('7.1. Thiết lập')
P('Máy Intel Core i5-9400F 2.90 GHz, Node.js 24.20.0, seed 20261017. Đồ thị dựng sẵn, tắt ghi trace, chỉ đo phần thuật toán; '
  'mỗi ô lấy trung vị 5 lần chạy (1 lần nếu quá 1,5 giây), có khởi động JIT trước. Hai họ đồ thị: **thưa** (m = 4n, n tới 10⁵) '
  'và **dày** (m ≈ n²/2, n tới 1000). Script: `scripts/benchmark.js`; số liệu: `4_Report/so-lieu/benchmark.md`, `.csv`.')
H2('7.2. Đồ thị thưa')
T(['n', 'm', 'Heap', 'Mảng', 'Bellman-Ford'], [
    ['1 000', '4 000', '2.2', '4.5', '2.6'], ['2 000', '8 000', '3.1', '18.1', '7.7'],
    ['5 000', '20 000', '6.5', '110.1', '25.8'], ['10 000', '40 000', '15.9', '435.2', '63.7'],
    ['20 000', '80 000', '31.8', '1513.7', '163.6'], ['50 000', '200 000', '97.8', '—', '365.0'],
    ['100 000', '400 000', '250.0', '—', '843.9'],
], widths=[2.6, 2.8, 3.2, 3.2, 3.2], center_cols=(0, 1, 2, 3, 4), title='Thời gian (ms) trên đồ thị thưa')
FIG('bd-thua.png', 'Thời gian chạy trên đồ thị thưa (thang log – log)', width=13.5)
P('Khi n tăng 10 lần (2 000 → 20 000), bản mảng chậm đi khoảng 84 lần (18.1 → 1513.7 ms), đúng bậc O(V²) (100 lần); '
  'bản heap chậm đi khoảng 10 lần (3.1 → 31.8 ms), đúng bậc O((V + E) log V). Ở n = 20 000, heap nhanh hơn mảng khoảng 48 lần. '
  'Số phép so sánh của bản mảng là n(n + 1)/2 (200 010 000 ở n = 20 000), trong khi heap chỉ có 26 853 lần push.')
H2('7.3. Đồ thị dày')
T(['n', 'm', 'Heap', 'Mảng', 'Bellman-Ford'], [
    ['100', '4 950', '0.4', '1.2', '2.1'], ['200', '19 900', '1.5', '4.2', '11.2'],
    ['400', '79 800', '9.0', '6.2', '63.9'], ['700', '244 650', '46.5', '26.9', '165.9'],
    ['1 000', '499 500', '97.3', '60.4', '550.4'],
], widths=[2.6, 2.8, 3.2, 3.2, 3.2], center_cols=(0, 1, 2, 3, 4), title='Thời gian (ms) trên đồ thị dày')
FIG('bd-day.png', 'Thời gian chạy trên đồ thị dày', width=13.5)
P('Trên đồ thị dày, từ khoảng n = 400 bản mảng vượt bản heap (ở n = 1000: 60.4 ms so với 97.3 ms), đúng như phân tích '
  'O(V²) so với O(V² log V) khi E ≈ V². Với n nhỏ hơn, bản heap vẫn nhanh hơn do hằng số nhỏ và nhiều phần tử cũ chưa tích lũy.')
H2('7.4. Nhận xét')
B('Kết quả đo khớp với lý thuyết: **heap cho đồ thị thưa, mảng cho đồ thị dày**.')
B('Bellman-Ford trên đồ thị ngẫu nhiên chỉ cần 6–15 vòng nhờ dừng sớm (thay vì n − 1), nên chậm hơn Dijkstra chỉ vài lần chứ không '
  'tới hàng nghìn lần như cận O(V · E). Ngược lại, trên test lưới `lon/04` nó cần nhiều vòng và mất khoảng 5 giây so với 0,24 giây của heap.')
B('Số phần tử cũ trong heap (lazy deletion) trên đồ thị thưa khoảng 34 % số đỉnh, trên đồ thị dày tới gần 5 lần số đỉnh; '
  'vẫn nhỏ so với E nên không ảnh hưởng bậc độ phức tạp.')
B('Hạn chế: đo trên một máy, JavaScript có JIT và bộ thu gom rác nên số đo dao động; vì vậy dùng trung vị và chỉ kết luận về xu hướng.')

# =====================================================================
# CHƯƠNG 8
# =====================================================================
H1('Quá trình xây dựng và sử dụng AI')
P('Chương này tường thuật lại quá trình làm đồ án, dựa trên `DEVLOG.md` (ghi sau mỗi yêu cầu, prompt chép nguyên văn) '
  'và lịch sử commit git. Mục tiêu là cho thấy rõ phần nào AI làm, phần nào sinh viên quyết định, và kết quả AI được kiểm chứng ra sao.')
H2('8.1. Công cụ và cách tổ chức làm việc')
T(['Công cụ', 'Model / phiên bản', 'Công đoạn'], [
    ['Claude trên claude.ai (Project môn học, có đính kèm slide và đề)', 'Claude Opus 5.5', 'Phân tích đề, lập kế hoạch, hỏi lý thuyết'],
    ['Claude Code (CLI, chạy trong thư mục dự án)', 'Claude Opus 5.5', 'Viết code, testcase, script kiểm thử, kịch bản video, bản nháp report'],
    ['Cytoscape.js', '3.30.2', 'Vẽ đồ thị trên giao diện'],
    ['Node.js', '24.20.0 LTS', 'Server tĩnh, testcase, benchmark'],
], widths=[7, 3.2, 5.8], size=11, title='Công cụ sử dụng')
P('Cách tổ chức: giai đoạn đầu dùng Claude trên claude.ai để phân tích đề và chốt phạm vi; kết quả được đúc kết thành hai file '
  'đặt ở thư mục dự án. `CLAUDE.md` ghi phạm vi, kiến trúc, cấu trúc trace, định dạng input/output và **quy tắc làm việc** '
  '(Claude Code tự đọc file này khi khởi động). `DEVLOG.md` là nhật ký. Quy tắc trong `CLAUDE.md` buộc Claude Code sau mỗi yêu cầu phải '
  'ghi một mục DEVLOG (prompt nguyên văn, tóm tắt, file sửa, vấn đề gặp phải), ghi lỗi vào "Nhật ký lỗi", commit git, và '
  '**để trống mục "Kiểm chứng của SV"** cho sinh viên tự viết.')

H2('8.2. Dòng thời gian')
TIMELINE = [
    ['#1', '11/10', 'Phân tích đề', '"Phân tích toàn bộ yếu cầu đố án sản phẩm cho môn này giúp tôi, dự tính của tôi làm mô phong thuật toán dijkstra"',
     'Bóc tách 4 sản phẩm, yêu cầu "thống nhất"; gợi ý so sánh nhiều chiến lược; 3 mức phạm vi, khuyến nghị mức B.'],
    ['#2', '13/10', 'Kế hoạch, lý thuyết', '"deadline là ngày 18/10" + câu hỏi về quy hoạch động và hành trình đường thư',
     'Giải thích Bellman-Ford/Floyd; 3 cách giải hành trình theo lịch; gợi ý `relax(edge, t)`; thu hẹp phạm vi. SV quyết định giữ bài toán đơn giản.'],
    ['#3', '14/10', 'Thiết lập', '"Giờ viết ra 1 file DEVLOG.md rồi tôi sẽ dùng claude cli ..."', 'Tạo `DEVLOG.md` và `CLAUDE.md`.'],
    ['#4', '14/10', 'Code lõi', '"Đọc CLAUDE.md và DEVLOG.md. Tạo khung thư mục ..., viết parser input, min-heap và Dijkstra (heap) sinh trace ... Chưa làm giao diện."',
     'Parser, graph, min-heap, Dijkstra heap, script in trace. Cài Node.js (hỏi SV trước). Phát hiện lỗi #1.'],
    ['#5–#6', '14/10', 'Tài liệu', '"tạo giúp tôi 1 file readme ..."; "git remote add origin ..."', 'README hướng dẫn chạy; gắn remote GitHub.'],
    ['#7–#8', '15/10', 'Giao diện', '"ủa hiện tại tôi phải chạy làm sao? sẽ dựng 1 trang web chứ nhỉ"; "có làm luôn đi, cập nhật trên readme luôn"',
     'Giao diện web 3 cột, phát lại trace, server tĩnh; kiểm tra trên Chrome.'],
    ['#9', '15/10', 'Sửa lỗi môi trường', '"chạy trên visual code bị:" + thông báo lỗi `npm.ps1`', 'Giải thích Execution Policy; hướng dẫn cách chạy khác; không tự đổi thiết lập bảo mật (lỗi #2).'],
    ['#10–#11', '15/10', 'Mở rộng', '"chúng ta còn báo cáo kịch bản demo và demo phải không?"; "có làm bước 1 luôn đi"',
     'Dijkstra mảng, Bellman-Ford, minh họa cạnh âm, chu trình âm; kiểm tra chéo 3000 đồ thị (lỗi #3, #4).'],
    ['#12', '17/10', 'Kiểm thử', '"làm tiếp tục bước tiếp theo đi"', '44 testcase, `run-tests.js`, benchmark, mutation test.'],
    ['#13', '17/10', 'Video', '"soạn giúp tôi kịch bản demo đi, hướng dẫn step by step, để tôi vừa quay màn hình vừa nói"', 'Kịch bản 8 cảnh lấy số bước từ trace thật.'],
    ['#14', '18/10', 'Report', '"vậy làm report giúp tôi đi, cái nào tôi dể review và chỉnh sửa thì làm cái đó"', 'Bản nháp report này (Word), ảnh chụp tự động, biểu đồ.'],
]
T(['#', 'Ngày', 'Công đoạn', 'Prompt (nguyên văn, có rút gọn bằng "...")', 'Kết quả'], TIMELINE,
  widths=[1.2, 1.3, 2.3, 5.6, 5.6], size=9.5, title='Các mốc chính (chi tiết đầy đủ trong DEVLOG.md)')
P('Lịch sử commit git (mỗi chức năng một commit, message tiếng Việt) là bằng chứng tiến trình: từ "Khung dự án, parser, '
  'min-heap, Dijkstra heap sinh trace" (14/10) tới "Thêm kịch bản video demo" (17/10). Kế hoạch ban đầu (mục #10) dự kiến '
  'làm testcase ngày 16/10; thực tế trễ một ngày.')

H2('8.3. Lỗi phát hiện và cách xử lý')
T(['#', 'Lỗi', 'Phát hiện bằng', 'Nguyên nhân → cách sửa'], [
    ['1', 'Đồ thị mẫu không sinh bước "bỏ qua vì đỉnh đã chốt" như comment khẳng định', 'Chạy script in trace, đọc trace', 'Chọn cạnh chưa đủ → thêm cạnh `6 4 1`'],
    ['2', '`npm` bị chặn trong terminal VS Code', 'SV chạy thử trên máy', 'PowerShell Execution Policy → chạy `node scripts/serve.js` hoặc `npm.cmd`'],
    ['3', 'Kiểm tra ngẫu nhiên gần như không sinh ca "cạnh âm, không chu trình âm" (1/3000)', 'Đếm số ca theo loại', 'Bộ sinh LCG lấy `% k` dùng bit thấp chu kỳ ngắn → đổi sang mulberry32 (387/3000)'],
    ['4', 'Ví dụ chu trình âm: các cạnh chồng lên nhau', 'Chụp màn hình Chrome headless', 'Bố cục cây xếp đỉnh thẳng hàng → bố cục vòng tròn'],
], widths=[0.8, 5.4, 3.8, 6], size=10.5, title='Nhật ký lỗi (từ DEVLOG)')
P('Lỗi #3 là ví dụ cho thấy kết quả "0 ca sai" có thể đánh lừa: phép kiểm tra chéo đầu tiên vẫn "đạt" nhưng gần như không '
  'kiểm tra trường hợp cạnh âm. Chỉ khi đếm số ca theo từng loại mới thấy bộ sinh ngẫu nhiên có vấn đề.')

H2('8.4. Kết quả của AI được kiểm chứng như thế nào')
P('Các kiểm chứng tự động đã thực hiện trong quá trình làm (ghi trong DEVLOG):')
B('Mọi thuật toán được đối chiếu với đáp án **độc lập** (vét cạn, Floyd-Warshall), không lấy Dijkstra làm đáp án cho chính nó.')
B('Mutation test (mục 6.5) để kiểm tra chính bộ test.')
B('Giao diện được kiểm tra bằng trình duyệt thật (Claude in Chrome, sau đó Chrome headless): đếm số bước, đọc output, kiểm tra console không lỗi.')
B('Số liệu trong report, kịch bản video được sinh từ trace và file kết quả thật, không chép tay.')
P('**Phần sinh viên tự kiểm chứng** (SV viết — ví dụ: chạy tay Dijkstra trên giấy với đồ thị 6 đỉnh và so với 23 bước '
  'trên giao diện; đọc lại code phần nào; giữ / sửa / bỏ đề xuất nào của AI và vì sao):')
P('[[SV viết đoạn này. Gợi ý: nội dung các mục "Kiểm chứng của SV" trong DEVLOG.]]', indent=False)

H2('8.5. Phân công giữa sinh viên và AI')
T(['Việc', 'Sinh viên', 'AI'], [
    ['Chọn đề, phạm vi, hạn nộp', 'Chọn Dijkstra; chốt hạn 18/10; chọn phạm vi; quyết định để hành trình đường thư làm hướng phát triển', 'Phân tích đề, đề xuất 3 mức phạm vi'],
    ['Quy ước dự án', '[[duyệt / sửa CLAUDE.md thế nào]]', 'Soạn `CLAUDE.md`, `DEVLOG.md`'],
    ['Mã nguồn', 'Ra yêu cầu từng bước, chạy thử trên máy, báo lỗi npm; [[đọc lại / sửa phần nào]]', 'Viết code, README, tự kiểm tra'],
    ['Testcase, benchmark', '[[...]]', 'Sinh test có seed, script chấm, benchmark'],
    ['Video', 'Quay màn hình và thuyết minh', 'Soạn kịch bản'],
    ['Report', '[[đọc lại, sửa, viết phần kiểm chứng và nhận xét]]', 'Bản nháp, ảnh chụp, biểu đồ'],
], widths=[3.4, 6.6, 6], size=10.5, title='Phân công')

H2('8.6. Nhận xét về việc dùng AI')
P('[[SV tự viết: AI giúp được gì, chỗ nào AI sai hoặc phải sửa, bài học rút ra. Một số sự kiện có thể dùng: AI tự phát hiện '
  'lỗi #1 và #3 khi kiểm tra; AI hỏi trước khi cài Node.js và không tự đổi Execution Policy; lịch làm bị trễ 1 ngày; '
  'quy tắc CLAUDE.md giúp DEVLOG được ghi đều.]]', indent=False)

# =====================================================================
# CHƯƠNG 9
# =====================================================================
H1('Kết luận và hướng phát triển')
H2('9.1. Kết quả đạt được')
B('Chương trình web mô phỏng từng bước Dijkstra (min-heap) với pseudo-code, hàng đợi, bảng dist/prev, đồ thị tô màu; tiến/lùi/tự chạy.')
B('Đủ phần mở rộng: Dijkstra mảng, Bellman-Ford, minh họa cạnh âm với bảng so sánh, phát hiện chu trình âm.')
B('44 testcase có đáp án độc lập, 128/128 lượt chạy đạt; mutation test cho thấy bộ test phát hiện được lỗi trong heap.')
B('Benchmark khớp lý thuyết: heap thắng trên đồ thị thưa (48 lần ở n = 20 000), mảng thắng trên đồ thị dày từ n ≈ 400.')
B('Report, video và testcase dùng chung một đồ thị; quá trình làm được ghi đầy đủ trong DEVLOG.')
H2('9.2. Hạn chế')
B('Giao diện chỉ vẽ tốt đồ thị ≤ 30 đỉnh.')
B('Ở bước "bỏ qua phần tử cũ", đỉnh bị bỏ qua vẫn tô màu "đang xét" (câu giải thích đã nói rõ).')
B('Chưa cài decrease-key / Fibonacci heap để so sánh với lazy deletion.')
B('Benchmark chỉ đo trên một máy.')
H2('9.3. Hướng phát triển')
B('**Hành trình đường thư theo lịch chuyến:** thay `relax(edge, t)` bằng hàm "giờ đến sớm nhất" theo bảng giờ chuyến; Dijkstra phụ thuộc thời gian vẫn đúng nhờ tính chất FIFO. Kiến trúc hiện tại đã chuẩn bị sẵn điểm mở rộng này.')
B('A* với heuristic khoảng cách, chạy trên bản đồ thật.')
B('So sánh decrease-key với lazy deletion; hiển thị heap dạng cây.')
B('Triển khai lên GitHub Pages để dùng trực tiếp không cần cài đặt.')

# =====================================================================
# TÀI LIỆU THAM KHẢO
# =====================================================================
H1_plain('Tài liệu tham khảo')
for i, ref in enumerate([
    'T. H. Cormen, C. E. Leiserson, R. L. Rivest, C. Stein, *Introduction to Algorithms*, 4th ed., MIT Press, 2022 — Chương 22 (Single-Source Shortest Paths).',
    'E. W. Dijkstra, "A note on two problems in connexion with graphs", *Numerische Mathematik*, vol. 1, pp. 269–271, 1959.',
    'R. Bellman, "On a routing problem", *Quarterly of Applied Mathematics*, vol. 16, no. 1, pp. 87–90, 1958.',
    'L. R. Ford Jr., "Network Flow Theory", RAND Corporation, P-923, 1956.',
    'M. Franz et al., "Cytoscape.js 2023 update: a graph theory library for visualization and analysis", *Bioinformatics*, vol. 39, no. 1, 2023.',
    '[[Slide bài giảng môn CS112 — tên tài liệu, giảng viên]]',
], 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    add_runs(p, f'[{i}]  {ref}')

# ---------------- Số trang ở chân trang (bỏ trang bìa) ----------------
sec.different_first_page_header_footer = True
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
r = fp.add_run()
it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'; r._r.append(it)
r = fp.add_run()
e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)

doc.core_properties.title = 'Báo cáo đồ án CS112 — Mô phỏng thuật toán Dijkstra'
doc.core_properties.author = '[Họ và tên SV]'
doc.save(OUT)
print('Đã ghi', OUT)
